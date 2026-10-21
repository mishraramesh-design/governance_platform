"""Document intake and strict contracts for the interactive testing workspace.

No document or model output is executable code. Results are derived from
browser assertions, never from a model's claim that a test passed.
"""
import hashlib
import ipaddress
import json
import re
import zipfile
from io import BytesIO
from urllib.parse import urlsplit

from pydantic import BaseModel, ConfigDict, Field, StringConstraints, field_validator
from typing import Annotated, Literal
from detectors import mask_sensitive_text

MAX_INPUT_BYTES = 2_000_000


def masked(value, limit=12000):
    return mask_sensitive_text(str(value or '')[:limit])[0]


def origin(url):
    if any(ord(char) < 32 for char in url) or '\\' in url:
        raise ValueError('URL control characters and backslashes are forbidden.')
    parsed = urlsplit(url)
    if parsed.scheme not in ('https', 'http') or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError('An HTTP(S) URL without credentials is required.')
    if parsed.port == 0:
        raise ValueError('Port zero is not allowed.')
    port = parsed.port or (443 if parsed.scheme == 'https' else 80)
    host = parsed.hostname.lower().rstrip('.')
    if ':' in host:
        if '%' in host:
            raise ValueError('Scoped IPv6 addresses are not supported.')
        ipaddress.ip_address(host)
        host = f'[{host}]'
    else:
        host = host.encode('idna').decode('ascii')
        if len(host) > 253 or not all(re.fullmatch(r'[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?', label) for label in host.split('.')):
            raise ValueError('Invalid hostname.')
    return f'{parsed.scheme}://{host}:{port}'


class Contract(BaseModel):
    model_config = ConfigDict(extra='forbid', str_strip_whitespace=True)


class Assertion(Contract):
    op: Literal['assert_text', 'assert_url']
    value: str = Field(min_length=1, max_length=1500)


class TestCase(Contract):
    id: str = Field(min_length=1, max_length=80, pattern=r'^[A-Za-z0-9_.-]+$')
    title: str = Field(min_length=3, max_length=300)
    source_ids: list[str] = Field(min_length=1, max_length=20)
    steps: list[str] = Field(min_length=1, max_length=15)
    expected: str = Field(min_length=3, max_length=1500)
    category: Literal['functional', 'non_functional'] = 'functional'
    oracle: Assertion | None = None
    use_case_id: str = Field(default='', max_length=80)
    execution: Literal['agent', 'page_assertion'] = 'agent'
    navigation_url: str = Field(default='', max_length=2048)

    @field_validator('steps')
    @classmethod
    def meaningful_steps(cls, values):
        steps = [item.strip() for item in values if item.strip()]
        if not steps or any(len(item) > 2000 for item in steps):
            raise ValueError('Provide nonempty test steps of at most 2000 characters each.')
        return steps


class UseCase(Contract):
    id: str = Field(min_length=1, max_length=80)
    title: str = Field(min_length=3, max_length=300)
    source_ids: list[str] = Field(min_length=1, max_length=20)
    actor: str = Field(default='Application user', max_length=200)
    preconditions: list[str] = Field(default_factory=list, max_length=10)
    main_flow: list[str] = Field(min_length=1, max_length=15)
    alternate_flows: list[str] = Field(default_factory=list, max_length=10)


class Plan(Contract):
    cases: list[TestCase] = Field(min_length=1, max_length=50)
    use_cases: list[UseCase] = Field(default_factory=list, max_length=50)

    @field_validator('cases')
    @classmethod
    def unique_cases(cls, cases):
        if len({case.id for case in cases}) != len(cases):
            raise ValueError('Test case IDs must be unique.')
        return cases


class BrowserAction(Contract):
    op: Literal['navigate', 'click', 'fill', 'select', 'press', 'scroll', 'assert_text', 'assert_url', 'handoff', 'blocked']
    node: str = Field(default='', max_length=80)
    value: Annotated[str, StringConstraints(strip_whitespace=False)] = Field(default='', max_length=2000)
    reason: str = Field(default='', max_length=1000)
    observation_id: str = Field(default='', max_length=80)


def json_object(text):
    text = text.strip()
    if text.startswith('```'):
        text = re.sub(r'^```(?:json)?\s*|\s*```$', '', text).strip()
    if len(text) > 100000:
        raise ValueError('Model output exceeds the workspace limit.')
    data = json.loads(text)
    if not isinstance(data, dict):
        raise ValueError('Expected one JSON object.')
    return data


def validate_plan(data, sources):
    plan = Plan.model_validate(data)
    known = {source['id'] for source in sources}
    use_case_ids = {item.id for item in plan.use_cases}
    if len(use_case_ids) != len(plan.use_cases):
        raise ValueError('Use case IDs must be unique.')
    for item in plan.use_cases:
        if not set(item.source_ids).issubset(known):
            raise ValueError(f'{item.id} references an unknown source.')
    for case in plan.cases:
        if not set(case.source_ids).issubset(known):
            raise ValueError(f'{case.id} references an unknown source.')
        if case.use_case_id and case.use_case_id not in use_case_ids:
            raise ValueError(f'{case.id} references an unknown use case.')
    return plan.model_dump()


def _check_zip(content):
    with zipfile.ZipFile(BytesIO(content)) as archive:
        if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 20_000_000:
            raise ValueError('Expanded document exceeds the 20 MB limit.')


def extract_document(filename, content, prefix):
    """Retain masked source statements and provenance; discard original upload."""
    if len(content) > MAX_INPUT_BYTES:
        raise ValueError('Each upload must be at most 2 MB.')
    extension = filename.lower().rsplit('.', 1)[-1]
    digest = hashlib.sha256(content).hexdigest()
    sources = []
    if extension == 'xlsx':
        _check_zip(content)
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in book.worksheets[:20]:
                headers = None
                for row in sheet.iter_rows(max_row=501, max_col=30):
                    values = ['[formula omitted]' if cell.data_type == 'f' else str(cell.value or '')[:2000] for cell in row]
                    if not any(values):
                        continue
                    if headers is None:
                        headers = values
                        continue
                    line = ' | '.join(f'{headers[i] or "Column " + str(i+1)}: {value}' for i, value in enumerate(values) if value)
                    sources.append({'id': f'{prefix}-S{book.worksheets.index(sheet)+1}-R{row[0].row}',
                                    'location': f'{sheet.title}!{row[0].row}', 'text': masked(line, 4000),
                                    'kind': 'excel_test_case', 'document_sha256': digest})
                    if len(sources) >= 300:
                        break
                if len(sources) >= 300:
                    break
        finally:
            book.close()
    else:
        if extension in ('txt', 'md', 'csv'):
            text = content.decode('utf-8-sig')
        elif extension == 'docx':
            _check_zip(content)
            from docx import Document
            doc = Document(BytesIO(content))
            text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
        elif extension == 'pdf':
            from pypdf import PdfReader
            doc = PdfReader(BytesIO(content))
            if len(doc.pages) > 200:
                raise ValueError('PDF documents are limited to 200 pages.')
            text = '\n'.join(p.extract_text() or '' for p in doc.pages)
        else:
            raise ValueError('Use XLSX, TXT, Markdown, CSV, DOCX or text-based PDF.')
        if len(text) > 500000:
            raise ValueError('Extracted text exceeds 500,000 characters.')
        for number, line in enumerate(text.splitlines(), 1):
            if len(line.strip()) < 5:
                continue
            sources.append({'id': f'{prefix}-L{number}', 'location': f'line {number}',
                            'text': masked(line, 2000), 'kind': 'requirement_candidate', 'document_sha256': digest})
            if len(sources) >= 300:
                break
    return sources, {'filename': masked(filename, 200), 'sha256': digest, 'source_count': len(sources),
                     'raw_retained': False, 'extraction_limit': 300}


def deterministic_plan(sources):
    """Reviewable fallback when the configured model is absent or over budget."""
    return {'cases': [{'id': f'TC-{i+1:03}', 'title': source['text'][:150],
                       'source_ids': [source['id']], 'steps': ['Review the linked source and define the browser workflow.'],
                       'expected': 'Owner must specify an observable assertion before this draft can be executed.',
                       'category': 'functional'} for i, source in enumerate(sources[:50])]}
