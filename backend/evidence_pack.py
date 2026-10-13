"""Evidence-pack generator: consolidates an audit (profile, findings, reports,
gate pack, API traceability and ledger hashes) into a DOCX or PDF for auditor
submission. Only masked evidence and hashes are included - never raw values."""
import io
from datetime import datetime, timezone

from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

EY_YELLOW = RGBColor(0xFF, 0xE6, 0x00)


def _summ(findings):
    by_sev, by_ctrl = {}, {}
    for f in findings:
        by_sev[f["severity"]] = by_sev.get(f["severity"], 0) + 1
        by_ctrl[f["control_name"]] = by_ctrl.get(f["control_name"], 0) + 1
    return by_sev, by_ctrl


def build_docx(data) -> bytes:
    d = Document()
    d.add_heading(f"IV&V Evidence Pack — {data['project']['name']}", 0)
    meta = d.add_paragraph()
    meta.add_run("Repository: ").bold = True
    meta.add_run(data["project"].get("repo_url", "—"))
    d.add_paragraph(f"Integrity level: {data['project'].get('integrity_level')}    "
                    f"Generated: {data['generated_at']}    Advisory only — observe & advise.")

    d.add_heading("1. Stack Profile", level=1)
    prof = data["profile"] or {}
    d.add_paragraph(f"Languages: {', '.join(f'{k} ({v})' for k, v in (prof.get('languages') or {}).items())}")
    d.add_paragraph(f"Frameworks: {', '.join(prof.get('frameworks') or []) or '—'}")
    d.add_paragraph(f"Scope: {prof.get('code_files_scanned', 0)} files, ~{prof.get('approx_loc', 0)} LOC")

    by_sev, by_ctrl = _summ(data["findings"])
    d.add_heading("2. Findings Summary", level=1)
    d.add_paragraph(f"Total findings: {len(data['findings'])}  |  By severity: " +
                    ", ".join(f"{k}: {v}" for k, v in by_sev.items()))
    t = d.add_table(rows=1, cols=2); t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text = "Control"; t.rows[0].cells[1].text = "Count"
    for k, v in by_ctrl.items():
        r = t.add_row().cells; r[0].text = k; r[1].text = str(v)

    d.add_heading("3. Top Findings", level=1)
    ft = d.add_table(rows=1, cols=5); ft.style = "Light Grid Accent 1"
    for i, h in enumerate(["Control", "Severity", "Verdict", "Subject", "L2 status"]):
        ft.rows[0].cells[i].text = h
    for f in sorted(data["findings"], key=lambda x: 0 if x["verdict"] == "FAIL" else 1)[:40]:
        r = ft.add_row().cells
        r[0].text = f["control_name"]; r[1].text = f["severity"]; r[2].text = f["verdict"]
        r[3].text = str(f.get("subject_ref", ""))[:48]; r[4].text = str(f.get("level2_status", ""))

    if data.get("endpoints"):
        d.add_heading("4. API Traceability", level=1)
        et = d.add_table(rows=1, cols=3); et.style = "Light Grid Accent 1"
        for i, h in enumerate(["Endpoint", "Defined in", "Tested by"]):
            et.rows[0].cells[i].text = h
        for e in data["endpoints"][:40]:
            r = et.add_row().cells
            r[0].text = f"{e.get('method')} {e.get('path')}"; r[1].text = str(e.get("defined_in", ""))[:40]
            r[2].text = (", ".join(e.get("tested_by", []))[:50] or "— untested —")

    d.add_heading("5. Reports", level=1)
    for rep in data["reports"]:
        d.add_heading(rep["title"], level=2)
        d.add_paragraph(rep.get("narrative", ""))
        for sec in rep.get("sections", []):
            d.add_paragraph(sec["heading"], style="Intense Quote")
            for b in sec.get("body", [])[:12]:
                d.add_paragraph(str(b), style="List Bullet")

    d.add_heading("6. Gate Pack", level=1)
    gt = d.add_table(rows=1, cols=3); gt.style = "Light Grid Accent 1"
    for i, h in enumerate(["Gate", "Status", "Reason"]):
        gt.rows[0].cells[i].text = h
    for g in data["gates"]:
        r = gt.add_row().cells
        r[0].text = f"{g['id']} {g['name']}"; r[1].text = g["status"]; r[2].text = g["reason"][:90]

    d.add_heading("7. Tamper-Evident Ledger", level=1)
    v = data["ledger_verify"]
    d.add_paragraph(f"Chain integrity: {'VALID' if v.get('valid') else 'BROKEN'}  |  "
                    f"Total entries: {v.get('total_entries')}  |  Last hash: {v.get('last_hash', '')[:32]}…")
    lt = d.add_table(rows=1, cols=4); lt.style = "Light Grid Accent 1"
    for i, h in enumerate(["Seq", "Type", "Ref", "Hash"]):
        lt.rows[0].cells[i].text = h
    for e in data["ledger_entries"][:60]:
        r = lt.add_row().cells
        r[0].text = str(e["seq"]); r[1].text = e["payload_type"]
        r[2].text = str(e["payload_ref"])[:24]; r[3].text = e["hash"][:24] + "…"

    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def build_pdf(data) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], textColor=colors.black)
    body = ss["BodyText"]
    el = []
    el.append(Paragraph(f"IV&V Evidence Pack — {data['project']['name']}", ss["Title"]))
    el.append(Paragraph(f"Repository: {data['project'].get('repo_url','—')}", body))
    el.append(Paragraph(f"Integrity level: {data['project'].get('integrity_level')} · Generated: {data['generated_at']} · Advisory only.", body))
    el.append(Spacer(1, 10))

    prof = data["profile"] or {}
    el.append(Paragraph("1. Stack Profile", h1))
    el.append(Paragraph(f"Languages: {', '.join(f'{k} ({v})' for k,v in (prof.get('languages') or {}).items())}", body))
    el.append(Paragraph(f"Frameworks: {', '.join(prof.get('frameworks') or []) or '—'}", body))
    el.append(Paragraph(f"Scope: {prof.get('code_files_scanned',0)} files, ~{prof.get('approx_loc',0)} LOC", body))
    el.append(Spacer(1, 8))

    by_sev, by_ctrl = _summ(data["findings"])
    el.append(Paragraph("2. Findings Summary", h1))
    el.append(Paragraph(f"Total: {len(data['findings'])} · " + ", ".join(f"{k}: {v}" for k, v in by_sev.items()), body))
    rows = [["Control", "Count"]] + [[k, str(v)] for k, v in by_ctrl.items()]
    el.append(_tbl(rows, [9 * cm, 3 * cm]))
    el.append(Spacer(1, 8))

    el.append(Paragraph("3. Top Findings", h1))
    frows = [["Control", "Sev", "Verdict", "Subject"]]
    for f in sorted(data["findings"], key=lambda x: 0 if x["verdict"] == "FAIL" else 1)[:30]:
        frows.append([f["control_name"][:26], f["severity"], f["verdict"], str(f.get("subject_ref", ""))[:30]])
    el.append(_tbl(frows, [6 * cm, 2 * cm, 2 * cm, 6 * cm]))
    el.append(Spacer(1, 8))

    if data.get("endpoints"):
        el.append(Paragraph("4. API Traceability", h1))
        erows = [["Endpoint", "Tested by"]]
        for e in data["endpoints"][:30]:
            erows.append([f"{e.get('method')} {e.get('path')}"[:40], (", ".join(e.get("tested_by", []))[:40] or "untested")])
        el.append(_tbl(erows, [8 * cm, 8 * cm]))
        el.append(Spacer(1, 8))

    el.append(Paragraph("5. Reports", h1))
    for rep in data["reports"]:
        el.append(Paragraph(rep["title"], ss["Heading2"]))
        el.append(Paragraph(rep.get("narrative", "").replace("\n", "<br/>"), body))
    el.append(Spacer(1, 8))

    el.append(Paragraph("6. Gate Pack", h1))
    grows = [["Gate", "Status"]] + [[f"{g['id']} {g['name']}", g["status"]] for g in data["gates"]]
    el.append(_tbl(grows, [11 * cm, 5 * cm]))
    el.append(Spacer(1, 8))

    el.append(Paragraph("7. Tamper-Evident Ledger", h1))
    v = data["ledger_verify"]
    el.append(Paragraph(f"Chain integrity: {'VALID' if v.get('valid') else 'BROKEN'} · entries: {v.get('total_entries')} · last hash: {v.get('last_hash','')[:40]}…", body))
    lrows = [["Seq", "Type", "Hash"]]
    for e in data["ledger_entries"][:40]:
        lrows.append([str(e["seq"]), e["payload_type"][:22], e["hash"][:30] + "…"])
    el.append(_tbl(lrows, [2 * cm, 6 * cm, 8 * cm]))

    doc.build(el)
    return buf.getvalue()


def _tbl(rows, widths):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.black),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#FFE600")),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
    ]))
    return t


def build(data, fmt):
    if fmt == "docx":
        return build_docx(data), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
    return build_pdf(data), "application/pdf", "pdf"
